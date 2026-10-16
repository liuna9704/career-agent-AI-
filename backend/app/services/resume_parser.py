import re
import unicodedata
from pathlib import Path

import fitz
from docx import Document


class ResumeParseError(ValueError):
    pass


def clean_text(text: str) -> str:
    text = unicodedata.normalize("NFKC", text)
    text = re.sub(r"[\x00-\x08\x0b\x0c\x0e-\x1f\x7f]", "", text)
    lines: list[str] = []
    previous = ""
    for raw_line in text.splitlines():
        line = re.sub(r"\s+", " ", raw_line).strip()
        line = re.sub(r"([^\w\s])\1{2,}", r"\1", line)
        if not line or line == previous:
            continue
        lines.append(line)
        previous = line
    return "\n".join(lines).strip()


def extract_text(file_path: str | Path) -> str:
    path = Path(file_path)
    suffix = path.suffix.lower()
    try:
        if suffix == ".pdf":
            with fitz.open(path) as document:
                raw_text = "\n".join(page.get_text("text") for page in document)
        elif suffix == ".docx":
            document = Document(path)
            paragraphs = [paragraph.text for paragraph in document.paragraphs]
            table_cells = [
                cell.text
                for table in document.tables
                for row in table.rows
                for cell in row.cells
            ]
            raw_text = "\n".join(paragraphs + table_cells)
        else:
            raise ResumeParseError("仅支持PDF或DOCX格式的简历")
    except ResumeParseError:
        raise
    except Exception as exc:
        raise ResumeParseError("简历文件损坏或内容无法读取") from exc

    result = clean_text(raw_text)
    if len(result) < 20:
        raise ResumeParseError("未能从简历中提取到足够文本，请确认文件不是扫描图片")
    return result

