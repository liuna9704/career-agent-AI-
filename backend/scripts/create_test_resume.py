from pathlib import Path

import fitz
from docx import Document

CONTENT = [
    "Zhang San",
    "Target Job: Java Developer",
    "Education: Bachelor of Computer Science",
    "Skills: Java, Spring Boot, MySQL, Redis, Docker",
    "Project: Intelligent Recruitment System",
    "Responsibility: Designed backend APIs and optimized database queries.",
    "Technology: Spring Boot, MySQL, Redis, Docker",
    "Result: Reduced average API latency by 35 percent.",
]


def create_fixtures() -> None:
    output_dir = Path(__file__).resolve().parents[1] / "tests" / "fixtures"
    output_dir.mkdir(parents=True, exist_ok=True)

    pdf = fitz.open()
    page = pdf.new_page()
    y = 72
    for line in CONTENT:
        page.insert_text((72, y), line, fontsize=11)
        y += 24
    pdf.save(output_dir / "test_resume.pdf")
    pdf.close()

    document = Document()
    document.add_heading("Test Resume", level=1)
    for line in CONTENT:
        document.add_paragraph(line)
    document.save(output_dir / "test_resume.docx")
    print(output_dir)


if __name__ == "__main__":
    create_fixtures()

